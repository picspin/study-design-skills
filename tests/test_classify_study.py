import importlib.util
import json
import subprocess
import tempfile
import unittest
from unittest import mock
from pathlib import Path


SCRIPT = Path(__file__).parents[1] / "skills/study-design-skills" / "scripts" / "classify_study.py"
SPEC = importlib.util.spec_from_file_location("classify_study", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)
SAMPLE_SCRIPT = SCRIPT.parent / "sample_size.py"
SAMPLE_SPEC = importlib.util.spec_from_file_location("sample_size", SAMPLE_SCRIPT)
SAMPLE_MODULE = importlib.util.module_from_spec(SAMPLE_SPEC)
SAMPLE_SPEC.loader.exec_module(SAMPLE_MODULE)
COMPILER_SCRIPT = SCRIPT.parent / "compile_study_spec.py"
COMPILER_SPEC = importlib.util.spec_from_file_location("compile_study_spec", COMPILER_SCRIPT)
COMPILER_MODULE = importlib.util.module_from_spec(COMPILER_SPEC)
COMPILER_SPEC.loader.exec_module(COMPILER_MODULE)
VALIDATOR_SCRIPT = SCRIPT.parent / "validate_study_spec.py"
VALIDATOR_SPEC = importlib.util.spec_from_file_location("validate_study_spec", VALIDATOR_SCRIPT)
VALIDATOR_MODULE = importlib.util.module_from_spec(VALIDATOR_SPEC)
VALIDATOR_SPEC.loader.exec_module(VALIDATOR_MODULE)
RUBRIC_SCRIPT = SCRIPT.parent / "select_rubrics.py"
RUBRIC_SPEC = importlib.util.spec_from_file_location("select_rubrics", RUBRIC_SCRIPT)
RUBRIC_MODULE = importlib.util.module_from_spec(RUBRIC_SPEC)
RUBRIC_SPEC.loader.exec_module(RUBRIC_MODULE)
EXTERNAL_SCRIPT = SCRIPT.parent / "external_evidence.py"
EXTERNAL_SPEC = importlib.util.spec_from_file_location("external_evidence", EXTERNAL_SCRIPT)
EXTERNAL_MODULE = importlib.util.module_from_spec(EXTERNAL_SPEC)
EXTERNAL_SPEC.loader.exec_module(EXTERNAL_MODULE)
DESIGN_SCRIPT = SCRIPT.parent / "design_study.py"
DESIGN_SPEC = importlib.util.spec_from_file_location("design_study", DESIGN_SCRIPT)
DESIGN_MODULE = importlib.util.module_from_spec(DESIGN_SPEC)
DESIGN_SPEC.loader.exec_module(DESIGN_MODULE)


class StudyTriageTests(unittest.TestCase):
    def test_jev_blend_is_conservative_and_shared_with_markdown(self):
        spec = {
            "study_title": "Survey", "study_type": "cross-sectional survey",
            "jev_review": {"status": "completed", "jev_score_10": 7.27},
        }
        _, rule_score, _, _, _, _ = DESIGN_MODULE.score_study_design(spec)
        combined = DESIGN_MODULE.combine_jev_score(rule_score, spec)
        self.assertLessEqual(combined, rule_score)
        self.assertIn(f"Overall score: {combined:.1f} / 10", DESIGN_MODULE.render_scoring_report(spec))

    def test_uncovered_jcr_category_is_not_replaced_by_general_medicine(self):
        import sys
        sys.path.insert(0, str(SCRIPT.parent))
        from generate_study_package import journal_recommendations, load_catalog

        journals, categories = journal_recommendations(
            {"target_jcr_category": "NURSING", "study_type": "cross-sectional survey"},
            load_catalog(),
        )
        self.assertEqual(categories, ["NURSING"])
        self.assertEqual(journals, [])

    def test_compiler_routes_legacy_rct_to_canonical_contract(self):
        result = COMPILER_MODULE.compile_spec({
            "study_title": "Pragmatic randomized trial",
            "study_type": "Randomized controlled trial",
            "primary_objective": "Compare 30-day adverse events between randomized groups",
            "population": "Eligible adults",
            "time_zero": "randomization",
            "groups": [{"label": "Intervention"}, {"label": "Control"}],
        })
        self.assertEqual(result["schema_version"], "1.0")
        self.assertEqual(result["confirmed_design"], "randomized_controlled_trial")
        self.assertEqual(result["route"]["family"], "randomized_trial")
        self.assertEqual(result["route"]["flow_layout"], "consort_trial")
        self.assertEqual(result["content_policy"]["table_cell_max_words"], 25)
        self.assertEqual(result["external_evidence"]["activation"], "disabled")

    def test_randomized_diagnostic_management_trial_uses_rct_route(self):
        spec = COMPILER_MODULE.compile_spec({
            "study_title": "Randomized diagnostic-management strategy trial",
            "study_type": "Pragmatic randomized diagnostic-management strategy trial",
            "confirmed_design": "randomized_controlled_trial",
            "primary_objective": "Compare management-plan change",
            "population": "Eligible adults",
            "time_zero": "randomization",
            "groups": [{"label": "Strategy A"}, {"label": "Strategy B"}],
            "reference_standard": "Composite lesion reference standard",
        })
        self.assertIn("CONSORT", DESIGN_MODULE.infer_guideline(spec))
        self.assertNotIn("one-gate", " ".join(DESIGN_MODULE.infer_design_warnings(spec)))
        self.assertIn("randomized allocation", " ".join(DESIGN_MODULE.render_matching(spec)))

    def test_external_evidence_is_default_disabled(self):
        self.assertEqual(EXTERNAL_MODULE.activation_decision(), {"activate": False, "reason": "default_disabled"})
        self.assertEqual(
            EXTERNAL_MODULE.activation_decision(classification_context_gap=True, allow_external_context=True),
            {"activate": True, "reason": "classification_context_gap"},
        )

    def test_springer_provider_normalizes_provenance(self):
        payload = {"result": [{"total": "1"}], "records": [{"title": "Example", "doi": "10.1000/example"}]}
        with mock.patch.dict("os.environ", {"NATURE_API_KEY": "test-key"}), mock.patch.object(
            EXTERNAL_MODULE, "_request_json", return_value=(200, {}, payload)
        ):
            result = EXTERNAL_MODULE.search_springer_open_access("keyword:trial", 1)
        self.assertEqual(result["provider"], "springer_open_access")
        self.assertEqual(result["total"], "1")
        self.assertNotIn("test-key", json.dumps(result))

    def test_product_specific_springer_key_precedes_fallback(self):
        payload = {"result": [{"total": "0"}], "records": []}
        with mock.patch.dict("os.environ", {"NATURE_META_API_KEY": "meta-key", "NATURE_API_KEY": "fallback-key"}), mock.patch.object(
            EXTERNAL_MODULE, "_request_json", return_value=(200, {}, payload)
        ) as request:
            EXTERNAL_MODULE.search_springer_meta("keyword:test", 1)
        self.assertEqual(request.call_args.kwargs["params"]["api_key"], "meta-key")

    def test_scopus_provider_uses_entitlement_warning(self):
        payload = {"search-results": {"opensearch:totalResults": "1", "entry": [{"dc:title": "Example"}]}}
        with mock.patch.dict("os.environ", {"SCOPUS_API_KEY": "test-key"}), mock.patch.object(
            EXTERNAL_MODULE, "_request_json", return_value=(200, {"X-RateLimit-Remaining": "19"}, payload)
        ):
            result = EXTERNAL_MODULE.search_scopus("TITLE-ABS-KEY(test)", 1)
        self.assertEqual(result["provider"], "scopus")
        self.assertIn("entitlements", " ".join(result["provenance"]["limitations"]))
        self.assertNotIn("test-key", json.dumps(result))

    def test_validator_flags_dense_clinical_narrative(self):
        result = VALIDATOR_MODULE.validate_spec({
            "study_title": "Cohort study",
            "study_type": "Observational cohort",
            "primary_objective": "Evaluate risk",
            "population": "Adults",
            "notes": ["First sentence. Second sentence. Third sentence. Fourth sentence."],
        })
        self.assertTrue(result["valid"])
        self.assertIn("density_note", {item["code"] for item in result["warnings"]})

    def test_compiler_preserves_external_provenance(self):
        result = COMPILER_MODULE.compile_spec({
            "study_title": "Cohort study",
            "study_type": "Observational cohort",
            "primary_objective": "Estimate event incidence",
            "population": "Eligible adults",
            "provenance": {
                "external_sources": [
                    {"source": "ClinicalTrials.gov", "retrieved": "2026-07-17"}
                ]
            },
        })
        self.assertEqual(result["provenance"]["external_sources"][0]["source"], "ClinicalTrials.gov")
        self.assertEqual(result["provenance"]["registry_version"], "1.0")

    def test_schema_validation_rejects_invalid_group_type(self):
        result = VALIDATOR_MODULE.validate_spec({
            "study_title": "Randomized trial",
            "study_type": "Randomized controlled trial",
            "primary_objective": "Compare event rates",
            "population": "Eligible adults",
            "time_zero": "randomization",
            "groups": "Intervention and control",
        })
        self.assertFalse(result["valid"])
        self.assertIn("schema", {item["code"] for item in result["errors"]})

    def test_route_selects_shared_and_design_specific_rubrics(self):
        result = RUBRIC_MODULE.select_rubrics({
            "study_title": "Diagnostic study",
            "study_type": "Diagnostic accuracy study",
            "primary_objective": "Estimate sensitivity and specificity",
            "population": "Consecutive suspected patients",
            "reference_standard": "Expert adjudication",
        })
        self.assertEqual(result["route"]["family"], "diagnostic_accuracy")
        self.assertEqual(result["rubric_files"], ["shared.csv", "diagnostic-accuracy.csv"])
        self.assertGreater(len(result["criteria"]), 10)

    def test_parallel_rct_sample_size_reports_analyzable_and_recruited(self):
        result = SAMPLE_MODULE.estimate_sample_size({"sample_size": {
            "method": "parallel_proportions",
            "control_event_rate": 0.25,
            "intervention_event_rate": 0.38,
            "alpha": 0.05,
            "power": 0.80,
            "loss_fraction": 0.15,
        }})
        self.assertEqual(result["status"], "estimated")
        self.assertGreater(result["recruited_n"], result["analyzable_n"])
        self.assertGreater(result["analyzable_n"], 350)

    def test_one_way_anova_sample_size_for_three_groups(self):
        result = SAMPLE_MODULE.estimate_sample_size({"sample_size": {
            "method": "one_way_anova", "effect_size_f": 0.25, "groups": 3,
            "alpha": 0.05, "power": 0.80, "loss_fraction": 0.10,
        }})
        self.assertEqual(result["status"], "estimated")
        self.assertEqual(result["analyzable_n"], 158)
        self.assertEqual(result["recruited_n"], 176)

    def test_correlation_sample_size_supports_multiplicity_alpha(self):
        result = SAMPLE_MODULE.estimate_sample_size({"sample_size": {
            "method": "correlation", "correlation": 0.30, "alpha": 0.01,
            "power": 0.80, "loss_fraction": 0.10,
        }})
        self.assertEqual(result["status"], "estimated")
        self.assertEqual(result["analyzable_n"], 125)
        self.assertEqual(result["recruited_n"], 139)

    def test_prevalence_precision_inflates_for_clustering_and_invalid_responses(self):
        result = SAMPLE_MODULE.estimate_sample_size({"sample_size": {
            "method": "prevalence_precision",
            "prevalence": 0.5,
            "half_width": 0.03,
            "alpha": 0.05,
            "design_effect": 2.0,
            "loss_fraction": 0.10,
        }})
        self.assertEqual(result["status"], "estimated")
        self.assertEqual(result["analyzable_n"], 2135)
        self.assertEqual(result["recruited_n"], 2373)

    def test_paired_sample_size_requires_directional_discordance(self):
        result = SAMPLE_MODULE.estimate_sample_size({"sample_size": {
            "method": "paired_binary",
            "discordant_control_only": 0.075,
            "discordant_intervention_only": 0.075,
        }})
        self.assertEqual(result["status"], "assumptions_required")

    def test_paired_imaging_does_not_trigger_ai_overlay(self):
        spec = {
            "proposal": "Prospective paired comparative imaging study with randomized order",
            "confirmed_design": "comparative_diagnostic_accuracy",
            "answers": {
                "primary_aim": "diagnostic_accuracy",
                "diagnostic_comparison": "paired",
                "data_source": "prospective_primary",
                "primary_outcome_family": "accuracy",
                "analysis_unit": "patient",
            },
        }

        result = MODULE.triage(spec)

        self.assertFalse(any("DECIDE-AI" in item for item in result["guideline_overlays"]))

    def test_asks_one_question_first(self):
        result = MODULE.triage({"proposal": "引入LLM智能体质控和监测预警系统，与传统工作流比较不良事件"})
        self.assertEqual(result["stage"], "clarify")
        self.assertEqual(result["next_question"]["id"], "rollout_structure")
        self.assertEqual(result["inferred_answers"]["primary_aim"], "intervention_effect")
        self.assertNotIn("questions", result)

    def test_retrospective_ehr_is_inferred_without_reasking(self):
        spec = {
            "proposal": "回顾性电子病历记录和影像学资料，评价LLM预警系统分阶段上线后的不良事件变化",
            "answers": {"rollout_structure": "nonrandom_staggered", "concurrent_control": "partial"},
        }
        result = MODULE.triage(spec)
        self.assertEqual(result["inferred_answers"]["data_source"], "retrospective_ehr")
        self.assertNotEqual(result["next_question"]["id"], "data_source")

    def test_llm_quality_control_converges_to_cits(self):
        spec = {
            "proposal": "引入LLM智能体质控和监测预警系统，与传统工作流比较不良事件",
            "answers": {
                "primary_aim": "intervention_effect",
                "rollout_structure": "fixed_date_series",
                "concurrent_control": "yes",
                "data_source": "retrospective_ehr",
                "primary_outcome_family": "clinical_safety",
                "analysis_unit": "time_period",
            },
        }
        result = MODULE.triage(spec)
        self.assertEqual(result["recommended_design"]["design_id"], "controlled_interrupted_time_series")
        self.assertIn("SQUIRE 2.0", result["guideline_overlays"])
        self.assertIn("TREND", result["guideline_overlays"])
        self.assertIn("RECORD", result["guideline_overlays"])
        self.assertEqual(result["stage"], "confirm")

    def test_prediction_uses_validation_not_psm(self):
        spec = {
            "proposal": "开发模型预测30天死亡风险",
            "answers": {
                "primary_aim": "prediction",
                "prediction_stage": "development",
                "data_source": "prospective_primary",
                "primary_outcome_family": "prediction_performance",
                "analysis_unit": "patient",
            },
            "confirmed_design": "prediction_model_development",
        }
        result = MODULE.triage(spec)
        self.assertEqual(result["stage"], "ready")
        methods = " ".join(result["recommended_design"]["methods"]).lower()
        self.assertIn("bootstrap", methods)
        self.assertNotIn("propensity", methods)

    def test_comparative_diagnostic_selects_quadas_c(self):
        spec = {
            "proposal": "比较两种影像AI的诊断准确性",
            "answers": {
                "primary_aim": "diagnostic_accuracy",
                "diagnostic_comparison": "paired",
                "data_source": "prospective_primary",
                "primary_outcome_family": "accuracy",
                "analysis_unit": "patient",
            },
            "confirmed_design": "comparative_diagnostic_accuracy",
        }
        result = MODULE.triage(spec)
        self.assertEqual(result["recommended_design"]["design_id"], "comparative_diagnostic_accuracy")
        self.assertIn("QUADAS-C", " ".join(result["recommended_design"]["bias_tools"]))

    def test_systematic_review_uses_characteristics_table(self):
        spec = {
            "proposal": "汇总多篇文献评价干预效果",
            "answers": {"primary_aim": "evidence_synthesis", "review_target": "intervention", "review_scope": "quantitative_meta", "data_source": "published_literature"},
            "confirmed_design": "systematic_review_meta_analysis",
        }
        result = MODULE.triage(spec)
        self.assertEqual(result["stage"], "ready")
        self.assertIn("study-characteristics", result["recommended_design"]["table"])

    def test_user_can_confirm_plausible_alternative(self):
        spec = {
            "proposal": "固定日期上线质控系统并观察多个前后时间点",
            "answers": {
                "primary_aim": "intervention_effect", "rollout_structure": "fixed_date_series",
                "concurrent_control": "no", "data_source": "retrospective_ehr",
                "primary_outcome_family": "clinical_quality", "analysis_unit": "time_period",
            },
            "confirmed_design": "interrupted_time_series",
        }
        result = MODULE.triage(spec)
        self.assertEqual(result["stage"], "ready")
        self.assertEqual(result["recommended_design"]["design_id"], "interrupted_time_series")

    def test_design_study_stops_before_confirmation(self):
        root = Path(__file__).parents[1]
        with tempfile.TemporaryDirectory() as temp_dir:
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/design_study.py"), str(root / "examples/llm_quality_proposal.json"), "--out-dir", temp_dir],
                check=True,
                capture_output=True,
                text=True,
            )
            files = {path.name for path in Path(temp_dir).iterdir()}
            self.assertEqual(files, {"study-intake.md", "study-intake.json"})
            result = json.loads((Path(temp_dir) / "study-intake.json").read_text(encoding="utf-8"))
            self.assertEqual(result["next_question"]["id"], "rollout_structure")

    def test_prediction_package_omits_smd(self):
        root = Path(__file__).parents[1]
        with tempfile.TemporaryDirectory() as temp_dir:
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/design_study.py"), str(root / "examples/prediction_model_confirmed.json"), "--out-dir", temp_dir, "--formats", "csv,md"],
                check=True,
                capture_output=True,
                text=True,
            )
            csv_path = next(Path(temp_dir).glob("*-table1.csv"))
            header = csv_path.read_text(encoding="utf-8-sig").splitlines()[0]
            self.assertNotIn("SMD", header)
            memo = next(Path(temp_dir).glob("*-memo.md")).read_text(encoding="utf-8")
            self.assertIn("PROBAST+AI", memo)
            self.assertIn("bootstrap", memo.lower())
            canonical = json.loads(next(Path(temp_dir).glob("*-study-package.json")).read_text(encoding="utf-8"))
            self.assertEqual(canonical["schema_version"], "1.0")
            self.assertEqual(canonical["route"]["family"], "prediction")

    def test_precomputed_table_rows_render_without_patient_level_data(self):
        root = Path(__file__).parents[1]
        spec = {
            "study_title": "Cross-sectional imaging biomarker study",
            "study_type": "Cross-sectional observational study",
            "confirmed_design": "descriptive_observational",
            "groups": [{"label": "Healthy", "n": 44}, {"label": "CKD", "n": 77}],
            "include_overall": False,
            "include_table_notes": True,
            "precomputed_table_rows": [{
                "characteristic": "Age, years",
                "groups": {"Healthy": "42.0 (12.0)", "CKD": "48.0 (11.0)"},
                "note": "Aggregate values supplied by investigators",
            }],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            spec_path = Path(temp_dir) / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/design_study.py"), str(spec_path), "--out-dir", temp_dir, "--formats", "csv"],
                check=True,
                capture_output=True,
                text=True,
            )
            table = next(Path(temp_dir).glob("*-table1.csv")).read_text(encoding="utf-8-sig")
            self.assertIn("Healthy (n=44)", table)
            self.assertIn("42.0 (12.0)", table)
            self.assertIn("Data-quality note", table)

    def test_confirmed_descriptive_imaging_route_uses_strobe_flow(self):
        root = Path(__file__).parents[1]
        spec = {
            "study_title": "Cross-sectional imaging biomarker study",
            "study_type": "Cross-sectional observational neuroimaging study",
            "confirmed_design": "descriptive_observational",
            "primary_objective": "Compare a continuous imaging biomarker across three groups",
            "population": "Adults undergoing research MRI",
            "groups": [{"label": "Control", "n": 40}, {"label": "Group A", "n": 30}, {"label": "Group B", "n": 30}],
            "variables": [{"name": "age", "label": "Age", "type": "continuous"}],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            spec_path = Path(temp_dir) / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/generate_study_package.py"), str(spec_path), "--out-dir", temp_dir, "--formats", "html,md"],
                check=True,
                capture_output=True,
                text=True,
            )
            report = next(Path(temp_dir).glob("*-report.html")).read_text(encoding="utf-8")
            memo = next(Path(temp_dir).glob("*-memo.md")).read_text(encoding="utf-8")
            self.assertIn("STROBE-style flow", report)
            self.assertNotIn("STARD-style participant and test flow", report)
            self.assertIn("Included cross-sectional sample", memo)
            self.assertNotIn("Reference standard performed", memo)

    def test_cross_sectional_survey_flow_uses_response_cleaning_nodes(self):
        root = Path(__file__).parents[1]
        spec = {
            "study_title": "National blood-culture practice survey",
            "study_type": "Cross-sectional questionnaire survey",
            "confirmed_design": "descriptive_observational",
            "primary_objective": "Estimate guideline-concordant collection practices",
            "population": "Nurses responding to a national online survey",
            "flow_counts": {
                "source_population": 45238,
                "excluded_before_cleaning": 2200,
                "cleaned_responses": 43038,
                "validity_discrepancy": 2913,
                "reported_valid": 40125,
                "regional_subset": 9436,
            },
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            spec_path = Path(temp_dir) / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/generate_study_package.py"), str(spec_path), "--out-dir", temp_dir, "--formats", "html,md"],
                check=True,
                capture_output=True,
                text=True,
            )
            report = next(Path(temp_dir).glob("*-report.html")).read_text(encoding="utf-8")
            memo = next(Path(temp_dir).glob("*-memo.md")).read_text(encoding="utf-8")
            self.assertIn("Questionnaires submitted", report)
            self.assertIn("Validity-status discrepancy", report)
            self.assertNotIn("MRI", report)
            self.assertIn("Cleaned response dataset", memo)
            self.assertNotIn("Measurements and prespecified quality control", memo)

    def test_survey_subcohort_flow_and_results_share_one_spec(self):
        root = Path(__file__).parents[1]
        spec = {
            "study_title": "Regional survey",
            "study_type": "Cross-sectional questionnaire survey",
            "confirmed_design": "descriptive_observational",
            "primary_objective": "Estimate practice prevalence",
            "population": "Regional survey respondents",
            "groups": [{"label": "Untrained", "n": 20}, {"label": "Trained", "n": 80}],
            "overall_n": 100,
            "precomputed_table_rows": [{"characteristic": "Hospital", "overall": "100", "groups": {"Untrained": "20", "Trained": "80"}}],
            "precomputed_result_rows": [{"Indicator": "Collection practice", "n/N": "79/100", "Percent": 79.0}],
            "flow_counts": {"source_population": 150, "outside_subcohort": 50, "subcohort_eligible": 100, "analysis_specific_excluded": 0, "primary_analysis": 100},
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            spec_path = Path(temp_dir) / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/generate_study_package.py"), str(spec_path), "--out-dir", temp_dir, "--formats", "xlsx,csv,html,md"],
                check=True, capture_output=True, text=True,
            )
            report = next(Path(temp_dir).glob("*-report.html")).read_text(encoding="utf-8")
            self.assertIn("Outside the prespecified subcohort", report)
            self.assertIn("Collection practice", report)
            self.assertNotIn("MRI", report)
            result_csv = next(Path(temp_dir).glob("*-results.csv")).read_text(encoding="utf-8-sig")
            self.assertIn("79/100", result_csv)
            from openpyxl import load_workbook
            workbook = load_workbook(next(Path(temp_dir).glob("*-package.xlsx")), read_only=True)
            self.assertIn("Results", workbook.sheetnames)
            workbook.close()

    def test_completed_aggregate_package_penalizes_missing_screening_and_imaging_qc(self):
        root = Path(__file__).parents[1]
        spec = {
            "study_title": "Completed cross-sectional imaging biomarker study",
            "study_type": "Cross-sectional observational neuroimaging study",
            "confirmed_design": "descriptive_observational",
            "clinical_area": "radiology",
            "target_jcr_category": "RADIOLOGY, NUCLEAR MEDICINE & MEDICAL IMAGING",
            "primary_objective": "Compare a continuous imaging biomarker",
            "population": "Adults undergoing research MRI",
            "groups": [{"label": "Control", "n": 40}, {"label": "CKD", "n": 60}],
            "variables": [{"name": "hypertension", "label": "Hypertension", "type": "categorical"}],
            "include_overall": False,
            "flow_counts": {"primary_analysis": 100},
            "precomputed_table_rows": [
                {"characteristic": "Hypertension", "groups": {"Control": "Not collected", "CKD": "40 (66.7)"}},
                {"characteristic": "Scanner/acquisition and motion QC", "groups": {"Control": "Not reported", "CKD": "Not reported"}},
            ],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            spec_path = Path(temp_dir) / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/generate_study_package.py"), str(spec_path), "--out-dir", temp_dir, "--formats", "html"],
                check=True,
                capture_output=True,
                text=True,
            )
            report = next(Path(temp_dir).glob("*-report.html")).read_text(encoding="utf-8")
            self.assertIn("lacks the assessed/screened denominator", report)
            self.assertIn("imaging-QC details are not reported", report)

    def test_systematic_review_package_uses_study_characteristics(self):
        root = Path(__file__).parents[1]
        with tempfile.TemporaryDirectory() as temp_dir:
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/design_study.py"), str(root / "examples/systematic_review_confirmed.json"), "--out-dir", temp_dir, "--formats", "xlsx,csv"],
                check=True,
                capture_output=True,
                text=True,
            )
            from openpyxl import load_workbook
            workbook = load_workbook(next(Path(temp_dir).glob("*-package.xlsx")), read_only=False)
            self.assertEqual(workbook.sheetnames[0], "Study Characteristics")
            workbook.close()
            header = next(Path(temp_dir).glob("*-table1.csv")).read_text(encoding="utf-8-sig").splitlines()[0]
            self.assertIn("Study,Year,Country/setting,Design", header)

    def test_exact_journal_target_outranks_non_ai_specialty_derivative(self):
        root = Path(__file__).parents[1]
        spec = {
            "study_title": "Paired MRI comparison",
            "study_type": "Comparative diagnostic accuracy study",
            "journal": "RADIOLOGY",
            "target_jcr_category": "RADIOLOGY, NUCLEAR MEDICINE & MEDICAL IMAGING",
            "groups": [{"label": "Sequence A"}, {"label": "Sequence B"}],
            "variables": [{"name": "age", "label": "Age", "type": "continuous"}],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            spec_path = Path(temp_dir) / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/design_study.py"), str(spec_path), "--out-dir", temp_dir, "--formats", "xlsx"],
                check=True,
                capture_output=True,
                text=True,
            )
            from openpyxl import load_workbook
            workbook = load_workbook(next(Path(temp_dir).glob("*-package.xlsx")), read_only=False)
            journal_sheet = workbook["Journal Fit"]
            self.assertEqual(journal_sheet["A5"].value, "RADIOLOGY")
            del journal_sheet
            workbook.close()

    def test_crossover_trial_uses_sequence_flow(self):
        root = Path(__file__).parents[1]
        spec = {
            "study_title": "Randomized technologist crossover study",
            "study_type": "Randomized controlled trial",
            "comparator": "Two devices within the same technologists",
            "groups": [{"label": "AB"}, {"label": "BA"}],
            "variables": [{"name": "experience", "label": "Experience", "type": "continuous"}],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            spec_path = Path(temp_dir) / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            result = subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/design_study.py"), str(spec_path)],
                check=True,
                capture_output=True,
                text=True,
            )
            self.assertIn("Technologists randomized to AB or BA device sequence", result.stdout)
            self.assertNotIn("Allocated to intervention", result.stdout)

    def test_html_uses_connected_crossover_enrollment_figure(self):
        root = Path(__file__).parents[1]
        spec = {
            "study_title": "Randomized technologist crossover study",
            "study_type": "Randomized controlled trial",
            "comparator": "Two devices within the same technologists",
            "groups": [{"label": "AB"}, {"label": "BA"}],
            "variables": [{"name": "experience", "label": "Experience", "type": "continuous"}],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            spec_path = Path(temp_dir) / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/design_study.py"), str(spec_path), "--out-dir", temp_dir, "--formats", "html"],
                check=True,
                capture_output=True,
                text=True,
            )
            report = next(Path(temp_dir).glob("*-report.html")).read_text(encoding="utf-8")
            self.assertIn('<figure class="enrollment-figure">', report)
            self.assertIn("CONSORT-style participant flow", report)
            self.assertIn("Sequence AB", report)
            self.assertIn("Excluded before randomization", report)
            self.assertIn("n = pending", report)
            self.assertNotIn('<ol class="flow">', report)

    def test_null_study_type_is_filled_from_confirmed_design(self):
        root = Path(__file__).parents[1]
        spec = {
            "proposal": "Randomize nursing technologists to AB or BA device sequences in a crossover trial",
            "confirmed_design": "randomized_controlled_trial",
            "study_type": None,
            "answers": {
                "primary_aim": "intervention_effect",
                "rollout_structure": "individual_randomized",
                "data_source": "prospective_primary",
                "primary_outcome_family": "clinical_quality",
                "analysis_unit": "clinician",
            },
            "groups": [{"label": "AB"}, {"label": "BA"}],
            "variables": [{"name": "experience", "label": "Experience", "type": "continuous"}],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            spec_path = Path(temp_dir) / "spec.json"
            spec_path.write_text(json.dumps(spec), encoding="utf-8")
            subprocess.run(
                ["python3", str(root / "skills/study-design-skills/scripts/design_study.py"), str(spec_path), "--out-dir", temp_dir, "--formats", "html"],
                check=True,
                capture_output=True,
                text=True,
            )
            report = next(Path(temp_dir).glob("*-report.html")).read_text(encoding="utf-8")
            self.assertIn("Randomized controlled trial", report)
            self.assertIn("CONSORT-style participant flow", report)


if __name__ == "__main__":
    unittest.main()
