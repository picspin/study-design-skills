#!/usr/bin/env python3
"""Generate publication-oriented Table 1 packages in CSV, XLSX, HTML, and Markdown."""

import argparse
import csv
import html
import json
import math
import re
from decimal import Decimal, ROUND_HALF_UP
from itertools import combinations
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from compile_study_spec import compile_spec
from design_study import (
    infer_design_warnings,
    infer_guideline,
    render,
    render_footnote,
    score_study_design,
)
from sample_size import estimate_sample_size

try:
    from scipy import stats
except ImportError:  # P values remain optional and are discouraged for baseline balance.
    stats = None


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_CATALOG = SCRIPT_DIR.parent / "references" / "jcr-2026-medicine-top69.csv"

CATEGORY_HINTS = {
    "cardiology": "CARDIAC & CARDIOVASCULAR SYSTEMS",
    "cardiovascular": "CARDIAC & CARDIOVASCULAR SYSTEMS",
    "oncology": "ONCOLOGY",
    "cancer": "ONCOLOGY",
    "radiology": "RADIOLOGY, NUCLEAR MEDICINE & MEDICAL IMAGING",
    "imaging": "RADIOLOGY, NUCLEAR MEDICINE & MEDICAL IMAGING",
    "neurology": "CLINICAL NEUROLOGY",
    "neuroscience": "NEUROSCIENCES",
    "gastroenterology": "GASTROENTEROLOGY & HEPATOLOGY",
    "hepatology": "GASTROENTEROLOGY & HEPATOLOGY",
    "urology": "UROLOGY & NEPHROLOGY",
    "nephrology": "UROLOGY & NEPHROLOGY",
    "kidney": "UROLOGY & NEPHROLOGY",
    "digital health": "MEDICAL INFORMATICS",
    "medical informatics": "MEDICAL INFORMATICS",
    "clinical ai": "MEDICAL INFORMATICS",
    "ai healthcare": "MEDICAL INFORMATICS",
}

GENERAL_MEDICINE = "MEDICINE, GENERAL & INTERNAL"
EXPERIMENTAL_MEDICINE = "MEDICINE, RESEARCH & EXPERIMENTAL"


def text_norm(value):
    return str(value or "").strip().casefold()


def include_smd(spec):
    if "include_smd" in spec:
        return bool(spec["include_smd"])
    route = spec.get("route") or {}
    if route.get("table_profile") == "exposure_balance":
        return True
    study = text_norm(spec.get("study_type"))
    return any(term in study for term in ["observational causal", "comparative effectiveness", "target trial", "real-world evidence", "rwe"])


def output_table_name(spec):
    table_profile = (spec.get("route") or {}).get("table_profile")
    if table_profile in {"study_characteristics", "evidence_map"}:
        return "Study Characteristics"
    if table_profile == "participant_context":
        return "Participant Context"
    study = text_norm(spec.get("study_type"))
    if any(term in study for term in ["systematic review", "meta-analysis"]):
        return "Study Characteristics"
    if "qualitative" in study:
        return "Participant Context"
    return "Table 1"


def p_value_note(spec):
    study = text_norm(spec.get("study_type"))
    if any(term in study for term in ["rct", "randomized", "stepped-wedge"]):
        return "Omitted; baseline significance testing does not validate randomization"
    if any(term in study for term in ["prediction", "prognostic"]):
        return "Omitted; emphasize events, calibration, discrimination, and validation"
    if any(term in study for term in ["diagnostic", "radiology", "imaging"]):
        return "Omitted from characteristics; accuracy comparisons require paired/design-specific inference"
    if any(term in study for term in ["interrupted time series", "difference-in-differences", "difference in differences"]):
        return "Omitted from characteristics; inference belongs to segmented/panel models"
    if any(term in study for term in ["systematic review", "meta-analysis", "scoping review"]):
        return "Not applicable to the study-characteristics table"
    if include_smd(spec):
        return "Omitted; use SMD for causal observational balance"
    return "Included only when a prespecified descriptive comparison requires it" if spec.get("include_p_values") else "Omitted by default"


def safe_slug(value):
    slug = re.sub(r"[^a-z0-9]+", "-", text_norm(value)).strip("-")
    return slug or "study-design"


def format_decimal(value, digits=1):
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    quantum = Decimal("1").scaleb(-digits)
    rounded = Decimal(str(float(value))).quantize(quantum, rounding=ROUND_HALF_UP)
    return f"{rounded:.{digits}f}"


def load_data(spec, spec_dir):
    source = spec.get("data_file")
    if not source:
        return None
    path = Path(source)
    if not path.is_absolute():
        path = spec_dir / path
    if path.suffix.casefold() in {".xlsx", ".xls"}:
        return pd.read_excel(path, sheet_name=spec.get("data_sheet", 0))
    if path.suffix.casefold() in {".csv", ".txt"}:
        return pd.read_csv(path)
    raise ValueError("data_file must be CSV, XLSX, or XLS")


def resolve_groups(spec, data):
    group_column = spec.get("group_column")
    configured = spec.get("groups") or []
    if data is None or not group_column:
        groups = []
        for index, group in enumerate(configured):
            if isinstance(group, dict):
                value = group.get("value", group.get("label", f"Group {index + 1}"))
                label = group.get("label", str(value))
                n = group.get("n", "")
            else:
                value = group
                label = str(group)
                n = ""
            groups.append({"value": value, "label": label, "n": n})
        return groups
    if group_column not in data.columns:
        raise ValueError(f"group_column '{group_column}' is not present in data_file")
    values = [value for value in pd.unique(data[group_column].dropna())]
    groups = []
    if configured:
        for group in configured:
            if isinstance(group, dict):
                value = group.get("value", group.get("label"))
                label = group.get("label", str(value))
            else:
                value = group
                label = str(group)
            groups.append({
                "value": value,
                "label": label,
                "n": int((data[group_column] == value).sum()),
            })
    else:
        for value in values:
            groups.append({
                "value": value,
                "label": str(value),
                "n": int((data[group_column] == value).sum()),
            })
    return groups


def resolve_variables(spec, data):
    variables = spec.get("variables") or []
    output = []
    if variables:
        for item in variables:
            if isinstance(item, str):
                item = {"name": item}
            variable = dict(item)
            variable.setdefault("label", variable.get("name", "Variable"))
            if data is not None and variable.get("name") in data.columns:
                inferred = "continuous" if pd.api.types.is_numeric_dtype(data[variable["name"]]) else "categorical"
            else:
                inferred = "categorical"
            variable.setdefault("type", inferred)
            variable.setdefault("summary", "mean_sd" if variable["type"] == "continuous" else "n_percent")
            output.append(variable)
        return output
    if data is None:
        defaults = [
            ("Age", "continuous", "mean_sd"),
            ("Sex", "categorical", "n_percent"),
            ("Race and ethnicity", "categorical", "n_percent"),
            ("Body mass index", "continuous", "mean_sd"),
            ("Disease severity", "categorical", "n_percent"),
            ("Key comorbidities", "categorical", "n_percent"),
            ("Prior medication use", "categorical", "n_percent"),
            ("Healthcare utilization", "continuous", "median_iqr"),
            ("Missing key covariates", "continuous", "n_missing"),
        ]
        return [{"name": name, "label": name, "type": kind, "summary": summary} for name, kind, summary in defaults]
    excluded = {spec.get("group_column"), spec.get("weight_column"), spec.get("id_column")}
    for name in [column for column in data.columns if column not in excluded][:25]:
        kind = "continuous" if pd.api.types.is_numeric_dtype(data[name]) else "categorical"
        output.append({
            "name": name,
            "label": name,
            "type": kind,
            "summary": "mean_sd" if kind == "continuous" else "n_percent",
        })
    return output


def weighted_mean_sd(series, weights=None):
    numeric = pd.to_numeric(series, errors="coerce")
    valid = numeric.notna()
    numeric = numeric[valid].astype(float)
    if numeric.empty:
        return math.nan, math.nan, 0
    if weights is None:
        return numeric.mean(), numeric.std(ddof=1), len(numeric)
    w = pd.to_numeric(weights[valid], errors="coerce").fillna(0).clip(lower=0).astype(float)
    keep = w > 0
    numeric, w = numeric[keep], w[keep]
    if numeric.empty or w.sum() <= 0:
        return math.nan, math.nan, 0
    mean = float(np.average(numeric, weights=w))
    variance = float(np.average((numeric - mean) ** 2, weights=w))
    return mean, math.sqrt(max(variance, 0)), len(numeric)


def format_continuous(series, summary, weights=None):
    numeric = pd.to_numeric(series, errors="coerce").dropna()
    if numeric.empty:
        return ""
    if summary == "median_iqr":
        q1, median, q3 = numeric.quantile([0.25, 0.5, 0.75])
        return f"{format_decimal(median)} ({format_decimal(q1)}-{format_decimal(q3)})"
    if summary == "n_missing":
        return str(int(series.isna().sum()))
    mean, sd, _ = weighted_mean_sd(series, weights)
    return f"{format_decimal(mean)} ({format_decimal(sd)})"


def continuous_smd(series_a, series_b, weights_a=None, weights_b=None):
    mean_a, sd_a, _ = weighted_mean_sd(series_a, weights_a)
    mean_b, sd_b, _ = weighted_mean_sd(series_b, weights_b)
    pooled = math.sqrt((sd_a ** 2 + sd_b ** 2) / 2) if not any(math.isnan(x) for x in [sd_a, sd_b]) else math.nan
    if not pooled or math.isnan(pooled):
        return math.nan
    return abs(mean_a - mean_b) / pooled


def binary_smd(p_a, p_b):
    pooled = math.sqrt((p_a * (1 - p_a) + p_b * (1 - p_b)) / 2)
    return abs(p_a - p_b) / pooled if pooled else 0.0


def weighted_rate(series, level, weights=None):
    valid = series.notna()
    if not valid.any():
        return math.nan, 0
    values = series[valid]
    if weights is None:
        return float((values == level).mean()), int((values == level).sum())
    w = pd.to_numeric(weights[valid], errors="coerce").fillna(0).clip(lower=0)
    total = float(w.sum())
    if total <= 0:
        return math.nan, 0
    return float(w[values == level].sum() / total), int((values == level).sum())


def p_value(variable, data, groups, group_column):
    if stats is None or len(groups) < 2:
        return ""
    samples = [data.loc[data[group_column] == group["value"], variable["name"]].dropna() for group in groups]
    try:
        if variable["type"] == "continuous":
            numeric = [pd.to_numeric(sample, errors="coerce").dropna() for sample in samples]
            result = stats.ttest_ind(*numeric, equal_var=False) if len(numeric) == 2 else stats.f_oneway(*numeric)
        else:
            table = pd.crosstab(data[variable["name"]], data[group_column])
            result = stats.chi2_contingency(table)
        value = float(result.pvalue if hasattr(result, "pvalue") else result[1])
        return "<0.001" if value < 0.001 else f"{value:.3f}"
    except (ValueError, TypeError):
        return ""


def table_columns(spec, data, groups):
    columns = ["Characteristic"]
    if spec.get("include_overall", True):
        n = len(data) if data is not None else spec.get("overall_n", "")
        columns.append(f"Overall (n={n})")
    columns.extend([f"{group['label']} (n={group['n']})" for group in groups])
    if include_smd(spec):
        columns.append("SMD")
    if spec.get("include_p_values", False):
        columns.append("P value")
    return columns


def build_table(spec, data, groups, variables):
    study = text_norm(spec.get("study_type"))
    if data is None and any(term in study for term in ["systematic review", "meta-analysis"]):
        columns = spec.get("characteristics_columns") or [
            "Study", "Year", "Country/setting", "Design", "Population", "N",
            "Intervention/exposure/test/model", "Comparator", "Primary outcome",
            "Follow-up", "Risk of bias",
        ]
        included = spec.get("included_studies") or []
        rows = [{column: study_row.get(column, "") for column in columns} for study_row in included]
        if not rows:
            rows = [{column: "" for column in columns}]
        return columns, rows
    columns = table_columns(spec, data, groups)
    if data is None:
        rows = []
        for variable in variables:
            row = {column: "" for column in columns}
            suffix = variable.get("unit")
            label = variable["label"] + (f", {suffix}" if suffix else "")
            row["Characteristic"] = label
            rows.append(row)
        return columns, rows

    group_column = spec.get("group_column")
    weight_column = spec.get("weight_column")
    overall_weights = data[weight_column] if weight_column and weight_column in data.columns else None
    rows = []
    for variable in variables:
        name = variable["name"]
        if name not in data.columns:
            continue
        label = variable["label"] + (f", {variable['unit']}" if variable.get("unit") else "")
        row = {column: "" for column in columns}
        row["Characteristic"] = label
        subsets = [data.loc[data[group_column] == group["value"]] for group in groups] if group_column else []
        if variable["type"] == "continuous":
            if spec.get("include_overall", True):
                row[columns[1]] = format_continuous(data[name], variable["summary"], overall_weights)
            for group, subset in zip(groups, subsets):
                weights = subset[weight_column] if weight_column and weight_column in subset.columns else None
                row[f"{group['label']} (n={group['n']})"] = format_continuous(subset[name], variable["summary"], weights)
            if include_smd(spec) and len(subsets) >= 2:
                smds = []
                for left, right in combinations(subsets, 2):
                    lw = left[weight_column] if weight_column and weight_column in left.columns else None
                    rw = right[weight_column] if weight_column and weight_column in right.columns else None
                    smds.append(continuous_smd(left[name], right[name], lw, rw))
                valid = [value for value in smds if not math.isnan(value)]
                row["SMD"] = format_decimal(max(valid), 3) if valid else ""
            if spec.get("include_p_values", False):
                row["P value"] = p_value(variable, data, groups, group_column)
            rows.append(row)
        else:
            levels = variable.get("levels") or [value for value in pd.unique(data[name].dropna())]
            max_smd = 0.0
            for left, right in combinations(subsets, 2):
                for level in levels:
                    left_rate, _ = weighted_rate(left[name], level, left[weight_column] if weight_column else None)
                    right_rate, _ = weighted_rate(right[name], level, right[weight_column] if weight_column else None)
                    if not math.isnan(left_rate) and not math.isnan(right_rate):
                        max_smd = max(max_smd, binary_smd(left_rate, right_rate))
            row["SMD"] = format_decimal(max_smd, 3) if include_smd(spec) and len(subsets) >= 2 else ""
            if spec.get("include_p_values", False):
                row["P value"] = p_value(variable, data, groups, group_column)
            rows.append(row)
            for level in levels:
                level_row = {column: "" for column in columns}
                level_row["Characteristic"] = f"  {level}"
                if spec.get("include_overall", True):
                    rate, count = weighted_rate(data[name], level, overall_weights)
                    level_row[columns[1]] = f"{count} ({format_decimal(rate * 100)}%)" if not math.isnan(rate) else ""
                for group, subset in zip(groups, subsets):
                    weights = subset[weight_column] if weight_column and weight_column in subset.columns else None
                    rate, count = weighted_rate(subset[name], level, weights)
                    level_row[f"{group['label']} (n={group['n']})"] = f"{count} ({format_decimal(rate * 100)}%)" if not math.isnan(rate) else ""
                rows.append(level_row)
        if spec.get("show_missing", True):
            missing = int(data[name].isna().sum())
            if missing:
                missing_row = {column: "" for column in columns}
                missing_row["Characteristic"] = "  Missing"
                if spec.get("include_overall", True):
                    missing_row[columns[1]] = f"{missing} ({format_decimal(missing / len(data) * 100)}%)"
                for group, subset in zip(groups, subsets):
                    count = int(subset[name].isna().sum())
                    missing_row[f"{group['label']} (n={group['n']})"] = f"{count} ({format_decimal(count / len(subset) * 100)}%)" if len(subset) else ""
                rows.append(missing_row)
    return columns, rows


def infer_target_categories(spec):
    explicit = spec.get("target_jcr_category")
    if explicit:
        return [str(explicit).upper()]
    haystack = " ".join([
        str(spec.get("clinical_area", "")),
        str(spec.get("study_type", "")),
        str(spec.get("journal", "")),
        " ".join(str(note) for note in (spec.get("notes") or [])),
    ]).casefold()
    categories = [category for hint, category in CATEGORY_HINTS.items() if hint in haystack]
    if not categories:
        categories = [GENERAL_MEDICINE, EXPERIMENTAL_MEDICINE]
    return list(dict.fromkeys(categories))


def load_catalog(path=DEFAULT_CATALOG):
    catalog = pd.read_csv(path)
    catalog["journal_key"] = catalog["journal"].str.casefold()
    return catalog


def journal_recommendations(spec, catalog, limit=8):
    categories = infer_target_categories(spec)
    requested = text_norm(spec.get("journal"))
    study_text = text_norm(spec.get("study_type"))
    is_ai = any(term in study_text for term in ["ai", "cdss", "triage", "workflow"])
    is_clinical = any(term in study_text for term in ["observational", "cohort", "rwe", "trial", "diagnostic", "clinical"])
    candidates = []
    for _, row in catalog.iterrows():
        category = str(row["wos_category"]).upper()
        name = str(row["journal"])
        score = 2.8
        reasons = []
        if category in categories:
            score += 4.5
            reasons.append("clinical-area category match")
        elif category in {GENERAL_MEDICINE, EXPERIMENTAL_MEDICINE}:
            score += 2.2
            reasons.append("broad medicine scope")
        if requested and requested == name.casefold():
            score += 2.0
            reasons.append("requested target journal")
        elif requested and requested in name.casefold():
            score += 0.4
            reasons.append("related journal-family title")
        if is_ai and any(token in name.casefold() for token in ["digital", "artificial intelligence", "image analysis"]):
            score += 1.5
            reasons.append("AI/digital-health fit")
        elif not is_ai and any(token in name.casefold() for token in ["artificial intelligence", "image analysis"]):
            score -= 1.5
            reasons.append("specialized AI/methods scope mismatch")
        if is_clinical and category in categories:
            score += 0.7
            reasons.append("clinical-design readership")
        impact_factor = float(row.get("impact_factor_2025", 0) or 0)
        score += min(impact_factor / 100, 0.8)
        candidates.append({
            "journal": name,
            "category": category,
            "impact_factor_2025": impact_factor,
            "fit_score": min(round(score, 1), 10.0),
            "rationale": "; ".join(reasons) or "reference-set candidate",
            "source_rank": int(row["source_rank"]),
        })
    frame = pd.DataFrame(candidates).sort_values(["fit_score", "impact_factor_2025"], ascending=False)
    frame = frame.drop_duplicates(subset=["journal"], keep="first").head(limit)
    return frame.to_dict("records"), categories


def category_benchmark_requirements(categories):
    requirements = [
        "Reconcile flowchart, Table 1, Methods population, and analysis denominators.",
        "Make the primary estimand, time anchor, missingness strategy, and sensitivity analyses explicit.",
    ]
    category_text = " ".join(categories)
    if "CARDIAC & CARDIOVASCULAR SYSTEMS" in category_text:
        requirements.extend([
            "Report established cardiovascular risk factors, disease severity, baseline therapies, procedures, and clinically relevant imaging or biomarker units.",
            "Prespecify cardiovascular endpoint definitions, adjudication, competing risks, and safety outcomes where applicable.",
        ])
    if "ONCOLOGY" in category_text:
        requirements.extend([
            "Report stage, performance status, line of therapy, biomarkers, metastatic burden, and prior treatment exposure.",
            "Separate intention-to-treat, safety, evaluable, and biomarker populations and define progression/death ascertainment.",
        ])
    if "RADIOLOGY" in category_text:
        requirements.extend([
            "Report patient/image/lesion/reader denominators, acquisition protocol, scanner/vendor, site, and reference standard.",
            "Address spectrum, verification, reference-standard, reader, device, and site-shift bias.",
        ])
    if "MEDICAL INFORMATICS" in category_text:
        requirements.extend([
            "Report workflow entry, AI trigger, output availability, clinician visibility, acceptance/override, downstream action, safety, and fairness.",
            "Separate model performance from clinical utility, adoption, implementation fidelity, and secular-trend effects.",
        ])
    if "GASTROENTEROLOGY & HEPATOLOGY" in category_text:
        requirements.append("Report disease etiology, severity/stage, decompensation or inflammatory activity, treatment history, and organ-specific outcomes.")
    if "CLINICAL NEUROLOGY" in category_text or "NEUROSCIENCES" in category_text:
        requirements.append("Report neurological phenotype, severity/functional scales, disease duration, treatment history, and ascertainment reliability.")
    if "UROLOGY & NEPHROLOGY" in category_text:
        requirements.append("Report renal/urologic disease stage, organ function, replacement therapy or procedure history, and competing-risk context.")
    if GENERAL_MEDICINE in category_text:
        requirements.append("Prioritize patient-important endpoints, broad generalizability, absolute effects, harms, and parsimonious main-text presentation.")
    return requirements


def flow_steps(spec):
    study = text_norm(spec.get("study_type"))
    counts = spec.get("flow_counts") or {}

    def node(label, key):
        value = counts.get(key)
        return f"{label} (n={value})" if value is not None else f"{label} (n not supplied)"

    if "controlled interrupted time series" in study or "cits" in study:
        return [
            node("Intervention and concurrent control source streams", "source_population"),
            node("Common eligibility and stable outcome definition", "eligible_at_time_zero"),
            node("Repeated pre-intervention observations by series", "pre_time_points"),
            node("Deployment date and ramp-up period", "intervention_point"),
            node("Repeated post-intervention observations by series", "post_time_points"),
            node("Controlled segmented-regression analysis", "primary_analysis"),
        ]
    if any(term in study for term in ["interrupted time series", "difference-in-differences", "difference in differences"]):
        primary = spec.get("primary_objective") or "Primary level/slope or group-by-time effect"
        secondary = spec.get("secondary_objectives") or ["Implementation, workflow, safety, subgroup, and robustness analyses"]
        if isinstance(secondary, str):
            secondary = [secondary]
        return [
            node("Clinical departments and source encounters", "source_population"),
            node("Stable eligibility, adverse-event definition, and denominator", "eligible_at_time_zero"),
            node("Repeated pre-rollout department-period observations", "pre_time_points"),
            node("Nonrandom rollout waves and prespecified ramp-up", "intervention_point"),
            node("Overlapping not-yet-treated/partial-control periods", "control_time_points"),
            node("Repeated post-rollout observations; excluded periods shown", "post_time_points"),
            node(f"Primary event-study/DiD analysis: {primary}", "primary_analysis"),
            node("Secondary analysis sets: " + "; ".join(str(item) for item in secondary), "secondary_analysis"),
        ]
    if any(term in study for term in ["observational", "cohort", "rwe", "real-world"]):
        return [
            node("Source population/data repository", "source_population"),
            node("Potentially eligible participants", "potentially_eligible"),
            node("Eligibility criteria applied before time zero", "eligible_at_time_zero"),
            node(f"Time zero/index date: {spec.get('time_zero', 'not specified')}", "time_zero"),
            node("Exposure/treatment groups assigned", "exposure_assigned"),
            node("Lookback and baseline-data sufficiency checked", "baseline_sufficient"),
            node("Matched/weighted analytic cohort", "designed_cohort"),
            node("Primary analysis cohort", "primary_analysis"),
            node("MI, complete-case, and alternative-design sensitivity cohorts", "sensitivity_analysis"),
        ]
    if any(term in study for term in ["rct", "randomized", "trial"]):
        primary = spec.get("primary_objective") or "Primary outcome and estimand assessed by randomized arm"
        secondary = spec.get("secondary_objectives") or ["Secondary effectiveness, diagnostic, safety, and follow-up analyses"]
        if isinstance(secondary, str):
            secondary = [secondary]
        crossover = "crossover" in text_norm(spec.get("study_title")) or "within the same" in text_norm(spec.get("comparator"))
        if crossover:
            return [
                node("Eligible nursing technologists and CT examinations", "source_population"),
                node("Technologists randomized to AB or BA device sequence", "randomized"),
                node("Period 1 device assignment; transition/training separated", "period_1"),
                node("Crossover to alternate device for Period 2", "period_2"),
                node("Image quality, scan success, repeats, workflow and harms ascertained", "outcomes_complete"),
                node(f"Primary mixed-effects ITT analysis: {primary}", "primary_analysis"),
                node("Period, sequence, learning and carryover sensitivity analyses", "sensitivity_analysis"),
                node("Secondary analysis sets: " + "; ".join(str(item) for item in secondary), "secondary_analysis"),
            ]
        return [
            node("Source population assessed for eligibility", "source_population"),
            node("Eligible and consented before randomization", "eligible_at_time_zero"),
            node(f"Time zero: {spec.get('time_zero', 'randomization')}", "time_zero"),
            node("Randomized with allocation concealment", "randomized"),
            node("Allocated intervention and control arms", "allocated"),
            node("Intervention/test completed; failures and crossovers retained", "completed_intervention"),
            node(f"Primary ITT analysis: {primary}", "primary_analysis"),
            node("Secondary analysis sets: " + "; ".join(str(item) for item in secondary), "secondary_analysis"),
        ]
    if any(term in study for term in ["diagnostic", "radiology", "imaging"]):
        primary = spec.get("primary_objective") or "Primary diagnostic-accuracy or paired-comparison analysis"
        secondary = spec.get("secondary_objectives") or ["Reader, lesion, subgroup, management-impact, and safety analyses"]
        if isinstance(secondary, str):
            secondary = [secondary]
        paired = "comparative" in study or "paired" in text_norm(spec.get("comparator"))
        test_step = "Both index tests completed in randomized order" if paired else "Index test completed"
        return [
            node("One-gate source population assessed", "source_population"),
            node("Eligible and consented participants", "eligible_at_time_zero"),
            node(test_step, "index_tests_complete"),
            node("Independent interpretation; indeterminate and failed tests retained", "interpretable_tests"),
            node("Common reference standard completed", "reference_standard_complete"),
            node("Complete paired/accuracy analysis set", "primary_analysis"),
            node(f"Primary analysis: {primary}", "primary_analysis"),
            node("Secondary analysis sets: " + "; ".join(str(item) for item in secondary), "secondary_analysis"),
        ]
    if any(term in study for term in ["ai", "cdss", "triage", "workflow"]):
        return ["Eligible encounters", "AI trigger", "Output generated", "Shown to clinician", "Accepted/modified/overridden", "Downstream action", "Safety and outcome ascertainment"]
    if any(term in study for term in ["systematic review", "meta-analysis"]):
        return [
            node("Records identified", "records_identified"),
            node("Duplicates removed", "records_deduplicated"),
            node("Titles/abstracts screened", "records_screened"),
            node("Full texts assessed", "full_texts"),
            node("Full texts excluded with reasons", "full_texts_excluded"),
            node("Studies included in qualitative synthesis", "included_studies"),
            node("Studies included in each meta-analysis", "meta_analysis_studies"),
        ]
    return ["Source population", "Eligibility", "Study groups", "Follow-up", "Final analysis"]


def enrollment_flow_html(spec):
    """Render a publication-style enrollment figure without external JS dependencies."""
    study = text_norm(spec.get("study_type"))
    layout = (spec.get("route") or {}).get("flow_layout", "")
    title_text = text_norm(f"{spec.get('study_title', '')} {spec.get('comparator', '')}")
    counts = spec.get("flow_counts") or {}
    exclusions = spec.get("flow_exclusions") or {}

    def count(key, prefix="n"):
        value = counts.get(key)
        return f"{prefix} = {html.escape(str(value))}" if value is not None else f"{prefix} = pending"

    def box(label, key=None, detail=None, kind="main"):
        number = f'<span class="flow-count">{count(key)}</span>' if key else ""
        extra = f'<span class="flow-detail">{html.escape(str(detail))}</span>' if detail else ""
        return f'<div class="flow-box {kind}"><strong>{html.escape(label)}</strong>{number}{extra}</div>'

    def exclusion(label, key, defaults):
        reasons = exclusions.get(key) or defaults
        if isinstance(reasons, str):
            reasons = [reasons]
        items = "".join(f"<li>{html.escape(str(item))}</li>" for item in reasons)
        return f'<div class="flow-box exclusion"><strong>{html.escape(label)}</strong><span class="flow-count">{count(key)}</span><ul>{items}</ul></div>'

    def down():
        return '<div class="flow-down" aria-hidden="true">&#8595;</div>'

    def side(main, excluded):
        return f'<div class="flow-row"><div>{main}</div><div class="flow-side-arrow" aria-hidden="true">&#8594;</div><div>{excluded}</div></div>'

    def branches(left, right, left_label=None, right_label=None):
        left_head = f'<div class="flow-branch-label">{html.escape(left_label)}</div>' if left_label else ""
        right_head = f'<div class="flow-branch-label">{html.escape(right_label)}</div>' if right_label else ""
        return f'<div class="flow-split" aria-hidden="true"><span></span><span></span></div><div class="flow-branches"><div>{left_head}{left}</div><div>{right_head}{right}</div></div>'

    def arm_side(main, excluded):
        return f'<div class="flow-arm-row"><div>{main}</div><div class="flow-arm-arrow" aria-hidden="true">&#8594;</div><div>{excluded}</div></div>'

    standard = "Study-specific participant flow"
    caption = "Flow of participants or examinations through eligibility, allocation/exposure, attrition, and final analysis."

    if layout == "prisma_review" or any(term in study for term in ["systematic review", "meta-analysis"]):
        standard = "PRISMA 2020"
        body = side(box("Records identified from databases and other sources", "records_identified"), exclusion("Records removed before screening", "records_removed", ["Duplicate records", "Automation or other prespecified removals"]))
        body += down() + side(box("Titles and abstracts screened", "records_screened"), exclusion("Records excluded", "records_excluded", ["Clearly ineligible by title/abstract"]))
        body += down() + side(box("Reports assessed for eligibility", "full_texts"), exclusion("Full-text reports excluded", "full_texts_excluded", ["Wrong population/design/outcome", "Insufficient or unavailable data", "Other prespecified reasons"]))
        body += down() + box("Studies included in qualitative synthesis", "included_studies", kind="final")
        body += down() + box("Studies included in each meta-analysis", "meta_analysis_studies", kind="final")
        caption = "PRISMA 2020 flow of records, reports, and studies through identification, screening, eligibility, and synthesis."
    elif layout in {"consort_trial", "stepped_wedge"} or any(term in study for term in ["rct", "randomized", "trial"]):
        standard = "CONSORT-style participant flow"
        crossover = "crossover" in title_text or "within the same" in title_text
        body = side(box("Assessed for eligibility", "source_population"), exclusion("Excluded before randomization", "excluded_before_randomization", ["Did not meet eligibility criteria", "Declined participation", "Other prespecified reasons"]))
        body += down() + box("Randomized", "randomized", kind="anchor")
        if crossover:
            left = box("Allocated to sequence AB", "sequence_ab") + down() + box("Period 1: information-enabled/new device", "ab_period_1") + down() + box("Period 2: comparator device", "ab_period_2") + down() + arm_side(box("Sequence AB retained", "ab_analyzed"), exclusion("Excluded", "ab_excluded", ["No evaluable examinations", "Protocol or outcome-data failure"]))
            right = box("Allocated to sequence BA", "sequence_ba") + down() + box("Period 1: comparator device", "ba_period_1") + down() + box("Period 2: information-enabled/new device", "ba_period_2") + down() + arm_side(box("Sequence BA retained", "ba_analyzed"), exclusion("Excluded", "ba_excluded", ["No evaluable examinations", "Protocol or outcome-data failure"]))
            body += branches(left, right, "Sequence AB", "Sequence BA")
            body += '<div class="flow-join" aria-hidden="true"><span></span><span></span></div>' + box("Included in mixed-effects intention-to-treat analysis", "primary_analysis", "Report technologist, period, and examination denominators", kind="final")
            caption = "CONSORT-style flow of nursing technologists and cardiovascular CT examinations through AB/BA sequence allocation, crossover periods, attrition, and analysis."
        else:
            left = box("Allocated to intervention", "intervention_allocated") + down() + box("Received intervention", "intervention_received") + down() + arm_side(box("Analyzed in intervention arm", "intervention_analyzed", kind="final"), exclusion("Lost/discontinued", "intervention_lost", ["Lost to follow-up", "Discontinued with reasons"]))
            right = box("Allocated to control", "control_allocated") + down() + box("Received control", "control_received") + down() + arm_side(box("Analyzed in control arm", "control_analyzed", kind="final"), exclusion("Lost/discontinued", "control_lost", ["Lost to follow-up", "Discontinued with reasons"]))
            body += branches(left, right, "Intervention", "Control")
            caption = "CONSORT-style participant flow through enrollment, randomization, allocation, follow-up, and analysis."
    elif layout in {"stard_accuracy", "stard_comparative"} or any(term in study for term in ["diagnostic", "radiology", "imaging"]):
        standard = "STARD-style participant and test flow"
        body = side(box("Patients/images assessed for eligibility", "source_population"), exclusion("Excluded before index testing", "excluded_before_index", ["Did not meet eligibility criteria", "Contraindication or unavailable imaging", "Other prespecified reasons"]))
        body += down() + side(box("Index test completed", "index_tests_complete"), exclusion("Index test unavailable or indeterminate", "index_test_excluded", ["Acquisition failure", "Non-evaluable or indeterminate result"]))
        body += down() + side(box("Reference standard completed", "reference_standard_complete"), exclusion("Reference standard unavailable", "reference_standard_excluded", ["Incomplete verification", "Timing or reference-standard criteria not met"]))
        body += down() + branches(box("Target condition present", "disease_present"), box("Target condition absent", "disease_absent"))
        body += '<div class="flow-join" aria-hidden="true"><span></span><span></span></div>' + box("Included in diagnostic accuracy analysis", "primary_analysis", "Report patient, image/lesion, and reader denominators separately", kind="final")
        caption = "STARD-style flow through eligibility, index testing, reference-standard verification, and diagnostic accuracy analysis."
    elif layout in {"tripod_development", "tripod_validation"} or any(term in study for term in ["prediction", "prognostic", "survival", "medical ai"]):
        standard = "TRIPOD+AI-style cohort flow"
        body = side(box("Source population/data repository", "source_population"), exclusion("Excluded before prediction time zero", "excluded_before_time_zero", ["Ineligible population or timing", "No usable outcome window", "Invalid or unavailable predictors"]))
        body += down() + box("Modeling cohort at prediction time zero", "modeling_cohort", "Report outcome events and censoring", kind="anchor")
        body += branches(box("Development and internal validation cohort", "development_cohort", "Include events"), box("External/temporal/geographic validation cohort", "external_validation_cohort", "Include events"), "Development", "Validation")
        caption = "TRIPOD+AI-style flow from source data through eligibility, prediction time zero, development, and independent validation."
    elif layout in {"controlled_time_series", "time_series", "comparative_panel"} or any(term in study for term in ["interrupted time series", "difference-in-differences", "difference in differences"]):
        standard = "STROBE/RECORD and SQUIRE-style encounter flow"
        body = side(box("Cardiovascular CT examinations in the source ward", "source_population"), exclusion("Excluded before cohort entry", "excluded_before_eligibility", ["Outside the prespecified CT population", "Duplicate/test/invalid injector record", "No linkable examination or outcome record"]))
        body += down() + box("Eligible examinations with stable definitions and denominators", "eligible_at_time_zero", kind="anchor")
        left = box("2025 pre-implementation period", "pre_period_exams", "Old-device injection protocol") + down() + arm_side(box("Retained pre-period observations", "pre_time_points"), exclusion("Excluded", "pre_period_excluded", ["Unapproved/unavailable records", "Non-comparable protocol or missing linkage"]))
        right = box("2026 post-implementation period", "post_period_exams", "Advanced new-device protocol bundle") + down() + arm_side(box("Retained post-period observations", "post_time_points"), exclusion("Excluded", "post_period_excluded", ["Prespecified ramp-up/transition window", "Non-comparable protocol or missing linkage"]))
        body += branches(left, right, "Pre-implementation", "Post-implementation")
        body += '<div class="flow-join" aria-hidden="true"><span></span><span></span></div>' + box("Included in segmented time-series analysis", "primary_analysis", "Report examinations, time points, events, and exposure denominators", kind="final")
        caption = "STROBE/RECORD and SQUIRE-style flow of cardiovascular CT examinations through eligibility, calendar-period assignment, exclusions, and segmented time-series analysis."
    elif any(term in study for term in ["ai", "cdss", "triage", "workflow"]):
        standard = "ICAML/DECIDE-AI-style clinical workflow flow"
        body = side(box("Eligible clinical encounters", "source_population"), exclusion("Excluded before AI workflow entry", "excluded_before_trigger", ["Outside deployment scope", "Missing required input", "Safety or governance exclusion"]))
        body += down() + side(box("AI trigger met and output generated", "ai_output_generated"), exclusion("AI output failure", "ai_output_failed", ["Unavailable data", "Timeout/interface failure", "Unsafe or invalid output"]))
        body += down() + side(box("Output shown to clinician/team", "ai_output_visible"), exclusion("Not reviewed", "not_reviewed", ["Workflow/staffing constraint", "Alert fatigue or interface issue"]))
        body += down() + box("Accepted, modified, overridden, or ignored", "clinician_action")
        body += down() + box("Outcome and safety ascertainment complete", "primary_analysis", kind="final")
        caption = "ICAML/DECIDE-AI-style flow from eligible encounters through AI triggering, clinician interaction, downstream action, and outcome ascertainment."
    else:
        standard = "STROBE/RECORD-style participant flow"
        body = side(box("Source population/data repository", "source_population"), exclusion("Excluded before eligibility", "excluded_before_eligibility", ["Outside sampling frame", "Duplicate or invalid record"]))
        body += down() + side(box("Potentially eligible participants", "potentially_eligible"), exclusion("Excluded before time zero", "excluded_before_time_zero", ["Eligibility criteria not met", "Insufficient baseline/lookback data"]))
        body += down() + box("Exposure groups assigned at time zero", "exposure_assigned", kind="anchor")
        body += down() + side(box("Follow-up and outcome ascertainment", "follow_up"), exclusion("Excluded/lost after cohort entry", "lost_after_entry", ["Loss to follow-up", "Outcome unavailable", "Report timing and bias implications"]))
        body += down() + box("Final primary analysis cohort", "primary_analysis", kind="final")

    return (
        f'<figure class="enrollment-figure"><div class="flow-standard">{html.escape(standard)}</div>'
        f'<div class="flow-canvas">{body}</div><figcaption><strong>Figure 1.</strong> {html.escape(caption)} '
        'Counts marked <em>pending</em> are protocol placeholders and must be reconciled with Table 1 before submission.</figcaption></figure>'
    )


def score_rows(spec, data=None):
    scores, total, ceiling, blockers, strengths, priorities = score_study_design(spec)
    scores = dict(scores)
    package_blockers = []
    package_cap = 10.0
    if data is not None and not spec.get("flow_counts"):
        package_blockers.append("Attrition-stage denominators were not supplied; the flowchart cannot be reconciled to Table 1.")
        score, maximum = scores["Enrollment, time anchor, and flowchart"]
        scores["Enrollment, time anchor, and flowchart"] = (max(0.0, round(score - 0.4, 1)), maximum)
        package_cap = min(package_cap, 7.0)
    matching_status = text_norm(spec.get("matching_status") or spec.get("analysis_stage"))
    if data is not None and spec.get("matching") and not any(term in matching_status for term in ["matched", "weighted", "completed", "post-match", "post match"]):
        package_blockers.append("Matching/weighting is planned but not documented as executed; Table 1 is treated as pre-design data.")
        score, maximum = scores["Statistical design and sensitivity analyses"]
        scores["Statistical design and sensitivity analyses"] = (max(0.0, round(score - 0.4, 1)), maximum)
        package_cap = min(package_cap, 7.5)
    imputation_status = text_norm(spec.get("imputation_status"))
    has_missing = data is not None and bool(data.isna().any().any())
    if has_missing and spec.get("missing_data") and not any(term in imputation_status for term in ["completed", "imputed", "pooled"]):
        package_blockers.append("Multiple imputation is planned but not documented as completed; displayed data retain missing values.")
        score, maximum = scores["Statistical design and sensitivity analyses"]
        scores["Statistical design and sensitivity analyses"] = (max(0.0, round(score - 0.3, 1)), maximum)
        package_cap = min(package_cap, 7.5)
    blockers = package_blockers + blockers
    priorities = package_blockers + priorities
    total = min(round(sum(score for score, _ in scores.values()), 1), package_cap, total)
    ceiling = max(ceiling, min(9.8, total + 2.0))
    rows = [{"domain": domain, "score": score, "maximum": maximum} for domain, (score, maximum) in scores.items()]
    return rows, total, ceiling, blockers, strengths, priorities


def write_csv(path, columns, rows):
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path, spec, columns, rows, journals, categories, scoring):
    workbook = Workbook()
    table_sheet = workbook.active
    table_sheet.title = output_table_name(spec)
    journal_sheet = workbook.create_sheet("Journal Fit")
    score_sheet = workbook.create_sheet("Benchmark")
    sample_sheet = workbook.create_sheet("Sample Size")
    notes_sheet = workbook.create_sheet("Methods Notes")

    dark = "162B36"
    accent = "0F766E"
    pale = "E8F3F1"
    warning = "FFF4D6"
    white = "FFFFFF"
    gray = "65747C"
    thin = Side(style="thin", color="D8E0E4")

    title = spec.get("study_title") or output_table_name(spec)
    table_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    table_sheet.cell(1, 1, title)
    title_size = 11 if len(title) > 105 else 12 if len(title) > 75 else 14
    table_sheet.cell(1, 1).font = Font(name="Arial", size=title_size, bold=True, color=white)
    table_sheet.cell(1, 1).fill = PatternFill("solid", fgColor=dark)
    table_sheet.cell(1, 1).alignment = Alignment(vertical="center", wrap_text=True)
    table_sheet.row_dimensions[1].height = 42 if len(title) > 75 else 28
    table_sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    subtitle = f"{spec.get('study_type', 'Study')} | {infer_guideline(spec.get('study_type'))}"
    table_sheet.cell(2, 1, subtitle)
    table_sheet.cell(2, 1).font = Font(name="Arial", size=9, color=gray)
    table_sheet.cell(2, 1).alignment = Alignment(vertical="center", wrap_text=True)
    table_sheet.row_dimensions[2].height = 28 if len(subtitle) > 90 else 20
    for col_index, column in enumerate(columns, start=1):
        cell = table_sheet.cell(4, col_index, column)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=accent)
        cell.alignment = Alignment(horizontal="left" if col_index == 1 else "center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=accent))
    first_column = columns[0]
    for row_index, row in enumerate(rows, start=5):
        is_level = str(row.get(first_column, "")).startswith("  ")
        for col_index, column in enumerate(columns, start=1):
            cell = table_sheet.cell(row_index, col_index, row.get(column, ""))
            cell.font = Font(name="Arial", size=9, bold=not is_level and col_index == 1)
            cell.alignment = Alignment(horizontal="left" if col_index == 1 else "right", vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if not is_level and col_index == 1:
                cell.fill = PatternFill("solid", fgColor="F3F6F7")
        if "SMD" in columns and row.get("SMD"):
            smd_cell = table_sheet.cell(row_index, columns.index("SMD") + 1)
            try:
                if float(row["SMD"]) >= 0.1:
                    smd_cell.fill = PatternFill("solid", fgColor=warning)
            except ValueError:
                pass
    footnote_row = 6 + len(rows)
    table_sheet.merge_cells(start_row=footnote_row, start_column=1, end_row=footnote_row + 2, end_column=len(columns))
    footnote = render_footnote(spec)
    if include_smd(spec):
        footnote += (
            " SMD is the maximum absolute pairwise standardized difference; for categorical variables it is the maximum binary-level SMD."
            " SMD <0.10 is a descriptive balance target, not proof of no confounding."
        )
    if spec.get("weight_column"):
        footnote += f" Summaries use nonnegative weights from {spec['weight_column']}; medians remain unweighted."
    table_sheet.cell(footnote_row, 1, footnote)
    table_sheet.cell(footnote_row, 1).font = Font(name="Arial", size=8, color=gray, italic=True)
    table_sheet.cell(footnote_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    table_sheet.freeze_panes = "B5"
    table_sheet.sheet_view.showGridLines = False
    table_sheet.column_dimensions["A"].width = 34
    for index in range(2, len(columns) + 1):
        table_sheet.column_dimensions[get_column_letter(index)].width = 20
    table_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    table_sheet.page_setup.fitToWidth = 1
    table_sheet.page_setup.fitToHeight = 0
    table_sheet.print_area = f"A1:{get_column_letter(len(columns))}{footnote_row + 2}"

    journal_sheet.append(["Recommended journals", "User-provided JCR 2026 reference set; metric column is 2025 impact factor"])
    journal_sheet.append(["Target categories", "; ".join(categories)])
    journal_sheet.append([])
    journal_headers = ["Journal", "WoS category", "2025 IF", "Scope fit /10", "Fit rationale"]
    journal_sheet.append(journal_headers)
    for journal in journals:
        journal_sheet.append([journal["journal"], journal["category"], journal["impact_factor_2025"], journal["fit_score"], journal["rationale"]])
    style_sheet(journal_sheet, len(journal_headers), dark, accent, white, thin)
    journal_sheet.column_dimensions["A"].width = 42
    journal_sheet.column_dimensions["B"].width = 44
    journal_sheet.column_dimensions["C"].width = 12
    journal_sheet.column_dimensions["D"].width = 15
    journal_sheet.column_dimensions["E"].width = 48

    score_rows_data, total, ceiling, blockers, strengths, priorities = scoring
    benchmark_family = " / ".join(item["journal"] for item in journals[:3])
    score_sheet.append(["JCR Top-Journal Benchmark", f"{total:.1f}/10", "Post-revision ceiling", f"{ceiling:.1f}/10"])
    score_sheet.append(["Domain", "Score", "Maximum", "Assessment"])
    for item in score_rows_data:
        score_sheet.append([item["domain"], item["score"], item["maximum"], "Heuristic evidence check; confirm manually against protocol and source data."])
    score_sheet.append([])
    score_sheet.append(["Benchmark family", benchmark_family])
    score_sheet.append(["Critical blockers", "; ".join(blockers) if blockers else "None automatically detected"])
    score_sheet.append(["Strengths", "; ".join(strengths) if strengths else "Requires manual assessment"])
    score_sheet.append(["Revision priorities", "; ".join(priorities[:5])])
    score_sheet.append(["Category requirements", " ".join(category_benchmark_requirements(categories))])
    for row_index in range(11, score_sheet.max_row + 1):
        score_sheet.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=4)
    style_sheet(score_sheet, 4, dark, accent, white, thin, header_row=2)
    score_sheet.column_dimensions["A"].width = 48
    score_sheet.column_dimensions["B"].width = 26
    score_sheet.column_dimensions["C"].width = 18
    score_sheet.column_dimensions["D"].width = 56

    sample = estimate_sample_size(spec)
    sample_sheet.append(["Sample Size Estimation", "Value"])
    sample_sheet.append(["Status", sample["status"]])
    sample_sheet.append(["Method", sample.get("method", "")])
    if sample["status"] == "estimated":
        sample_sheet.append(["Analyzable sample", f"{sample['analyzable_n']} {sample['unit']}"])
        sample_sheet.append(["Recruitment target", f"{sample['recruited_n']} {sample['recruited_unit']}"])
        sample_sheet.append(["Formula", sample["formula"]])
        for key, value in sample["assumptions"].items():
            sample_sheet.append([f"Assumption: {key}", value])
    else:
        for item in sample.get("missing", []):
            sample_sheet.append(["Required input", item])
    for item in sample.get("caveats", []):
        sample_sheet.append(["Caveat", item])
    style_sheet(sample_sheet, 2, dark, accent, white, thin, header_row=1)
    sample_sheet.column_dimensions["A"].width = 34
    sample_sheet.column_dimensions["B"].width = 110

    notes = [
        ["Confirmed design", spec.get("confirmed_design", "Not triaged")],
        ["Study type", spec.get("study_type", "")],
        ["Guideline", infer_guideline(spec.get("study_type"))],
        ["Guideline stack", "; ".join(spec.get("guideline_stack", []))],
        ["Bias/appraisal tools", "; ".join(spec.get("bias_tools", []))],
        ["Analysis/validation methods", "; ".join(spec.get("analysis_methods", []))],
        ["Time zero/index anchor", spec.get("time_zero", "Not specified")],
        ["Group column", spec.get("group_column", "Not specified")],
        ["Weight column", spec.get("weight_column", "Not used")],
        ["P values", p_value_note(spec)],
        ["JCR source", "JCR-70.xlsx supplied by user; 69 records, 68 unique journal titles"],
        ["Category benchmark", " ".join(category_benchmark_requirements(categories))],
        ["Interpretation", "Journal scope-fit and benchmark scores are editorial/methodological aids, not acceptance probabilities."],
    ]
    notes_sheet.append(["Methods and provenance", "Value"])
    for item in notes:
        notes_sheet.append(item)
    for warning_text in infer_design_warnings(spec):
        notes_sheet.append(["Design warning", warning_text])
    style_sheet(notes_sheet, 2, dark, accent, white, thin, header_row=1)
    notes_sheet.column_dimensions["A"].width = 28
    notes_sheet.column_dimensions["B"].width = 110

    workbook.save(path)
    workbook.close()


def style_sheet(sheet, columns, dark, accent, white, thin, header_row=4):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f"A{header_row + 1}"
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=12, bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=dark)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    if sheet.max_row >= header_row:
        for cell in sheet[header_row]:
            cell.font = Font(name="Arial", size=9, bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=accent)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in sheet.iter_rows(min_row=header_row + 1, max_col=columns):
        for cell in row:
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)


def html_table(columns, rows):
    head = "".join(f"<th>{html.escape(str(column))}</th>" for column in columns)
    body = []
    first_column = columns[0]
    for row in rows:
        label = str(row.get(first_column, ""))
        level = label.startswith("  ")
        cells = []
        for column in columns:
            value = row.get(column, "")
            class_name = "characteristic" if column == first_column else "numeric"
            cells.append(f'<td class="{class_name}">{html.escape(str(value))}</td>')
        body.append(f'<tr class="{"level" if level else "variable"}">{"".join(cells)}</tr>')
    return f'<div class="table-wrap"><table><thead><tr>{head}</tr></thead><tbody>{"".join(body)}</tbody></table></div>'


def write_html(path, spec, columns, rows, journals, categories, scoring):
    score_rows_data, total, ceiling, blockers, strengths, priorities = scoring
    journal_rows = "".join(
        f"<tr><td>{html.escape(item['journal'])}</td><td>{html.escape(item['category'])}</td>"
        f"<td>{item['impact_factor_2025']:.1f}</td><td><strong>{item['fit_score']:.1f}</strong></td>"
        f"<td>{html.escape(item['rationale'])}</td></tr>" for item in journals
    )
    domain_rows = "".join(
        f"<tr><td>{html.escape(item['domain'])}</td><td>{item['score']:.1f}/{item['maximum']:.1f}</td></tr>" for item in score_rows_data
    )
    enrollment_figure = enrollment_flow_html(spec)
    blocker_html = "".join(f"<li>{html.escape(item)}</li>" for item in blockers) or "<li>No automatic critical blocker detected; manual review remains required.</li>"
    priority_html = "".join(f"<li>{html.escape(item)}</li>" for item in priorities[:5])
    warnings_html = "".join(f"<li>{html.escape(item)}</li>" for item in infer_design_warnings(spec))
    requirements_html = "".join(f"<li>{html.escape(item)}</li>" for item in category_benchmark_requirements(categories))
    benchmark_family = " / ".join(item["journal"] for item in journals[:3])
    sample = estimate_sample_size(spec)
    if sample["status"] == "estimated":
        assumption_rows = "".join(f"<tr><td>{html.escape(str(key))}</td><td>{html.escape(str(value))}</td></tr>" for key, value in sample["assumptions"].items())
        sample_html = (
            f'<div class="scores"><div class="panel"><div class="label">Analyzable sample</div><div class="big">{sample["analyzable_n"]}</div><div>{html.escape(sample["unit"])}</div></div>'
            f'<div class="panel"><div class="label">Recruitment target</div><div class="big">{sample["recruited_n"]}</div><div>{html.escape(sample["recruited_unit"])}</div></div></div>'
            f'<div class="panel"><strong>{html.escape(sample["method"])}</strong><p><code>{html.escape(sample["formula"])}</code></p></div>'
            f'<div class="table-wrap"><table><thead><tr><th>Assumption</th><th>Value</th></tr></thead><tbody>{assumption_rows}</tbody></table></div>'
        )
    else:
        required = "".join(f"<li>{html.escape(item)}</li>" for item in sample.get("missing", []))
        sample_html = f'<div class="panel warning"><strong>Sample size not calculated</strong><p>Method: {html.escape(sample.get("method", "Not specified"))}</p><ul>{required}</ul></div>'
    sample_caveats = "".join(f"<li>{html.escape(item)}</li>" for item in sample.get("caveats", []))
    document = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(spec.get('study_title') or 'Study design report')}</title>
<style>
:root{{--ink:#162b36;--teal:#0f766e;--paper:#f6f8f8;--line:#d8e0e4;--muted:#65747c;--warn:#a83b24}}
*{{box-sizing:border-box}} body{{margin:0;background:var(--paper);color:var(--ink);font:15px/1.55 Arial,sans-serif}}
header{{background:var(--ink);color:white;padding:42px max(24px,calc((100vw - 1180px)/2)) 34px;border-bottom:6px solid var(--teal)}}
header p{{max-width:900px;color:#cbd7db;margin:8px 0 0}} h1{{font-size:32px;line-height:1.15;margin:0;letter-spacing:0}}
main{{max-width:1180px;margin:0 auto;padding:28px 24px 60px}} h2{{font-size:21px;margin:34px 0 12px;border-bottom:2px solid var(--ink);padding-bottom:7px}}
.meta,.scores{{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:12px}} .panel{{background:white;border:1px solid var(--line);border-radius:6px;padding:17px}}
.label{{color:var(--muted);font-size:12px;text-transform:uppercase}} .big{{font-size:30px;font-weight:700;color:var(--teal)}}
.table-wrap{{overflow:auto;background:white;border:1px solid var(--line)}} table{{width:100%;border-collapse:collapse;font-size:13px}} th{{background:var(--teal);color:white;text-align:left;padding:10px;position:sticky;top:0}} td{{padding:8px 10px;border-bottom:1px solid var(--line)}} td.numeric{{text-align:right;white-space:nowrap}} tr.variable td:first-child{{font-weight:700;background:#f1f5f5}} tr.level td:first-child{{padding-left:26px}}
	.enrollment-figure{{margin:0;background:white;border:1px solid var(--line);padding:20px;overflow:auto}} .flow-standard{{font-size:12px;font-weight:700;color:var(--teal);text-transform:uppercase;margin-bottom:16px}}
	.flow-canvas{{max-width:900px;min-width:620px;margin:0 auto;padding:4px 18px 8px}} .flow-box{{border:2px solid #285b69;background:#fff;padding:11px 14px;text-align:center;min-height:66px;display:flex;flex-direction:column;justify-content:center;box-shadow:0 1px 0 rgba(22,43,54,.08)}}
	.flow-box.anchor{{background:#e2f0f1}} .flow-box.final{{background:#285b69;color:#fff}} .flow-box.exclusion{{border-width:1.5px;text-align:left;background:#f8faf9}} .flow-box.exclusion ul{{margin:5px 0 0;padding-left:18px;font-size:12px;line-height:1.35}}
	.flow-count{{display:block;color:#0b62c4;font-weight:700;margin-top:3px}} .flow-box.final .flow-count{{color:#fff}} .flow-detail{{display:block;font-size:12px;color:var(--muted);margin-top:4px}} .flow-box.final .flow-detail{{color:#e0ecee}}
	.flow-down{{height:34px;text-align:center;font-size:29px;line-height:34px;color:#285b69;font-weight:400}} .flow-row{{display:grid;grid-template-columns:minmax(0,1fr) 52px minmax(0,1fr);align-items:center}} .flow-side-arrow{{font-size:29px;text-align:center;color:#285b69}}
	.flow-arm-row{{display:grid;grid-template-columns:minmax(0,1fr) 24px minmax(0,1fr);align-items:center}} .flow-arm-row .flow-box{{padding:9px;min-height:92px;font-size:12px}} .flow-arm-row .flow-box.exclusion ul{{font-size:10px;padding-left:14px}} .flow-arm-arrow{{text-align:center;color:#285b69;font-size:20px}}
	.flow-branches{{display:grid;grid-template-columns:1fr 1fr;gap:48px}} .flow-branch-label{{text-align:center;font-size:12px;font-weight:700;color:var(--teal);text-transform:uppercase;margin:0 0 7px}} .flow-split,.flow-join{{height:42px;position:relative;margin:0 24%}}
	.flow-split::before{{content:"";position:absolute;left:0;right:0;top:20px;border-top:2px solid #285b69}} .flow-split::after{{content:"";position:absolute;left:50%;top:0;height:21px;border-left:2px solid #285b69}} .flow-split span::before{{content:"";position:absolute;top:20px;height:15px;border-left:2px solid #285b69}} .flow-split span::after{{content:"";position:absolute;top:33px;width:0;height:0;border-left:6px solid transparent;border-right:6px solid transparent;border-top:8px solid #285b69}} .flow-split span:first-child::before{{left:0}} .flow-split span:last-child::before{{right:0}} .flow-split span:first-child::after{{left:-5px}} .flow-split span:last-child::after{{right:-5px}}
	.flow-join::before{{content:"";position:absolute;left:0;right:0;top:0;border-top:2px solid #285b69}} .flow-join::after{{content:"";position:absolute;left:calc(50% - 5px);top:22px;width:0;height:0;border-left:6px solid transparent;border-right:6px solid transparent;border-top:8px solid #285b69}} .flow-join span:first-child::before,.flow-join span:last-child::before{{content:"";position:absolute;top:0;height:21px;border-left:2px solid #285b69}} .flow-join span:first-child::before{{left:0}} .flow-join span:last-child::before{{right:0}} .flow-join span:first-child::after{{content:"";position:absolute;left:50%;top:0;height:23px;border-left:2px solid #285b69}} figcaption{{max-width:900px;margin:15px auto 0;color:var(--muted);font-size:12px}}
	.warning{{border-left:4px solid var(--warn)}} .fine{{color:var(--muted);font-size:12px}} ul,ol{{padding-left:20px}} @media print{{header{{padding:24px}} main{{padding:16px}} .table-wrap{{overflow:visible}}}}
	@media (max-width:720px){{.enrollment-figure{{padding:12px}} .flow-canvas{{transform-origin:top left}}}}
</style></head><body>
<header><h1>{html.escape(spec.get('study_title') or 'Study Design and Table 1 Report')}</h1><p>{html.escape(spec.get('population', 'Study population'))}</p></header>
<main>
<section class="meta"><div class="panel"><div class="label">Study design</div><strong>{html.escape(str(spec.get('study_type','Unspecified')))}</strong></div><div class="panel"><div class="label">Guideline</div><strong>{html.escape(infer_guideline(spec.get('study_type')))}</strong></div><div class="panel"><div class="label">Time zero</div><strong>{html.escape(str(spec.get('time_zero','Not specified')))}</strong></div><div class="panel"><div class="label">Target categories</div><strong>{html.escape('; '.join(categories))}</strong></div></section>
<h2>{html.escape(output_table_name(spec))}</h2>{html_table(columns, rows)}<p class="fine">The characteristics object, denominator rules, balance metrics, and inferential columns are selected from the confirmed study design.</p>
	<h2>Enrollment and analysis flow</h2>{enrollment_figure}
<h2>Sample size estimation</h2>{sample_html}<div class="panel"><strong>Interpretation caveats</strong><ul>{sample_caveats}</ul></div>
<h2>Journal scope fit</h2><div class="table-wrap"><table><thead><tr><th>Journal</th><th>WoS category</th><th>2025 IF</th><th>Scope fit /10</th><th>Rationale</th></tr></thead><tbody>{journal_rows}</tbody></table></div><p class="fine">Source: user-provided JCR-70.xlsx, treated as a JCR 2026 reference set with a 2025 impact-factor column. Scope-fit scores are not acceptance probabilities and do not substitute for current author-instruction checks.</p>
<h2>JCR top-journal benchmark</h2><p><strong>Benchmark family:</strong> {html.escape(benchmark_family)}</p><div class="scores"><div class="panel"><div class="label">Current score</div><div class="big">{total:.1f}/10</div></div><div class="panel"><div class="label">Post-revision ceiling</div><div class="big">{ceiling:.1f}/10</div></div></div><div class="table-wrap"><table><thead><tr><th>Domain</th><th>Score</th></tr></thead><tbody>{domain_rows}</tbody></table></div>
<div class="meta"><div class="panel warning"><strong>Critical blockers</strong><ul>{blocker_html}</ul></div><div class="panel"><strong>Revision priorities</strong><ol>{priority_html}</ol></div></div>
<h2>Category benchmark requirements</h2><div class="panel"><ul>{requirements_html}</ul></div>
<h2>Design-specific method stack</h2><div class="meta"><div class="panel"><strong>Reporting guidelines</strong><ul>{''.join(f'<li>{html.escape(item)}</li>' for item in spec.get('guideline_stack', []))}</ul></div><div class="panel"><strong>Bias/appraisal tools</strong><ul>{''.join(f'<li>{html.escape(item)}</li>' for item in spec.get('bias_tools', []))}</ul></div><div class="panel"><strong>Analysis/validation methods</strong><ul>{''.join(f'<li>{html.escape(item)}</li>' for item in spec.get('analysis_methods', []))}</ul></div></div>
<h2>Design warnings</h2><div class="panel warning"><ul>{warnings_html}</ul></div>
</main></body></html>"""
    path.write_text(document, encoding="utf-8")


def generate_package(spec, spec_dir, output_dir, formats):
    spec = compile_spec(spec)
    output_dir.mkdir(parents=True, exist_ok=True)
    data = load_data(spec, spec_dir)
    groups = resolve_groups(spec, data)
    variables = resolve_variables(spec, data)
    columns, rows = build_table(spec, data, groups, variables)
    catalog_path = Path(spec.get("journal_catalog", DEFAULT_CATALOG))
    if not catalog_path.is_absolute():
        catalog_path = spec_dir / catalog_path
    catalog = load_catalog(catalog_path)
    journals, categories = journal_recommendations(spec, catalog)
    scoring = score_rows(spec, data)
    stem = safe_slug(spec.get("output_name") or spec.get("study_title") or "study-design")
    canonical_path = output_dir / f"{stem}-study-package.json"
    canonical_path.write_text(json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
    outputs = [canonical_path]
    if "csv" in formats:
        path = output_dir / f"{stem}-table1.csv"
        write_csv(path, columns, rows)
        outputs.append(path)
    if "xlsx" in formats:
        path = output_dir / f"{stem}-package.xlsx"
        write_xlsx(path, spec, columns, rows, journals, categories, scoring)
        outputs.append(path)
    if "html" in formats:
        path = output_dir / f"{stem}-report.html"
        write_html(path, spec, columns, rows, journals, categories, scoring)
        outputs.append(path)
    if "md" in formats or "markdown" in formats:
        path = output_dir / f"{stem}-memo.md"
        path.write_text(render(spec), encoding="utf-8")
        outputs.append(path)
    manifest = {
        "study_title": spec.get("study_title"),
        "data_rows": len(data) if data is not None else None,
        "table_rows": len(rows),
        "journal_reference_records": len(catalog),
        "unique_journals": int(catalog["journal_key"].nunique()),
        "schema_version": spec["schema_version"],
        "route": spec["route"],
        "outputs": [str(path) for path in outputs],
    }
    manifest_path = output_dir / f"{stem}-manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    outputs.append(manifest_path)
    return outputs


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("spec", type=Path, help="Path to JSON study specification")
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--formats", default="xlsx,csv,html,md", help="Comma-separated: xlsx,csv,html,md")
    args = parser.parse_args()
    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    formats = {item.strip().casefold() for item in args.formats.split(",") if item.strip()}
    supported = {"xlsx", "csv", "html", "md", "markdown"}
    invalid = formats - supported
    if invalid:
        raise ValueError(f"Unsupported formats: {', '.join(sorted(invalid))}")
    for output in generate_package(spec, args.spec.parent, args.out_dir, formats):
        print(output)


if __name__ == "__main__":
    main()
